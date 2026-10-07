from openpyxl import load_workbook
from html import escape
from pathlib import Path


# =========================
# 文件位置
# =========================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "keypoint.xlsx"
OUTPUT_FILE = BASE_DIR / "index.html"


# =========================
# 读取 Excel
# =========================

wb = load_workbook(EXCEL_FILE, data_only=True)
ws = wb.active


items = []

for row in ws.iter_rows(min_row=2, values_only=True):

    # Excel:
    # A = ID
    # B = 答案
    # C = 前半段
    # D = 後半段
    # E = 分類

    item_id, answer, front, back, category = row[:5]

    # ID 为空的行跳过
    if item_id is None:
        continue

    item_id = str(item_id).strip()
    answer = "" if answer is None else str(answer)
    front = "" if front is None else str(front)
    back = "" if back is None else str(back)
    category = "" if category is None else str(category)

    items.append({
        "id": item_id,
        "answer": answer,
        "front": front,
        "back": back,
        "category": category
    })


# =========================
# 按分类分组
# =========================

categories = {}

for item in items:

    category = item["category"]

    if category not in categories:
        categories[category] = []

    categories[category].append(item)


# =========================
# 生成知识点 HTML
# =========================

content = []

for category, category_items in categories.items():

    if category:
        content.append(
            f'<h2>{escape(category)}</h2>'
        )

    content.append(
        '<ul class="knowledge-list">'
    )

    for item in category_items:

        item_id = escape(item["id"])
        answer = escape(item["answer"])
        front = escape(item["front"])
        back = escape(item["back"])

        content.append(f'''
<li class="item" data-id="{item_id}">

    <input
        class="item-check"
        type="checkbox"
    >

    <span class="front">
        {front}
    </span>

    <span
        class="answer"
        tabindex="0"
        role="button"
    >
        {answer}
    </span>

    <span class="back">
        {back}
    </span>

</li>
''')

    content.append('</ul>')


knowledge_html = "\n".join(content)


# =========================
# 生成完整 index.html
# =========================

html = f'''<!DOCTYPE html>
<html lang="ja">

<head>

<meta charset="UTF-8">

<meta
    name="viewport"
    content="width=device-width, initial-scale=1.0"
>

<title>宅建知識</title>

</head>


<body>

<main>

<h1>宅建知識</h1>

{knowledge_html}

</main>


<!-- 加载你原来的 JavaScript -->
<script src="https://www.gstatic.com/firebasejs/10.12.5/firebase-app-compat.js"></script>
<script src="https://www.gstatic.com/firebasejs/10.12.5/firebase-database-compat.js"></script>

<script src="app.js"></script>

</body>

</html>
'''


# =========================
# 写入 index.html
# =========================

OUTPUT_FILE.write_text(
    html,
    encoding="utf-8"
)


# =========================
# 完成提示
# =========================

print("================================")
print("HTML生成完成！")
print("================================")
print(f"Excel：{EXCEL_FILE.name}")
print(f"HTML ：{OUTPUT_FILE.name}")
print(f"知识点数量：{len(items)}")
print("================================")
